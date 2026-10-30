"""data/derived -> outputs/ 제출용 엑셀·CSV

읍면 8개를 한 행씩 놓은 결과표를 만든다. derived 만 읽고 raw 는 읽지 않는다.
종합 점수는 계산하지 않는다. '종합점수' 열은 그 사실을 남기려고 비워 둔다.

  python tools/export_excel.py
"""
from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parents[1]
DERIVED = ROOT / "data" / "derived"
REGISTRY = ROOT / "data" / "sources.json"
OUT = ROOT / "outputs"
BASENAME = "군위군_주거환경취약지역_기초지표_2025"

ANALYSIS_PERIOD = "2025-12"      # 인구 분모 기준 시점. 시설 자료 기준일과 가장 가깝다
SIGUNGU = "군위군"

# (컬럼명, 자료형, 단위, 설명) - 결과표 열 순서와 컬럼설명 시트를 한 곳에서 정한다
COLUMNS = [
    ("종합점수", "실수", "점", "계산하지 않는다. 근거 있는 가중치가 없어 비워 둔다 (docs/methodology.md §3)"),
    ("행정구역코드", "문자열", "", "행정안전부 행정기관코드 10자리. 조인키. 정수로 바꾸면 조인이 깨진다"),
    ("시군구명", "문자열", "", "군 이름"),
    ("읍면명", "문자열", "", "읍·면 이름"),
    ("총인구", "정수", "명", "2025년 12월 말 주민등록 인구 (행정안전부)"),
    ("고령인구", "정수", "명", "65세 이상 주민등록 인구"),
    ("고령화율", "실수", "%", "고령인구 / 총인구 × 100"),
    ("노령화지수", "실수", "", "유소년(0~14세) 100명당 고령인구. SGIS 2024년 (등록센서스)"),
    ("노년부양비", "실수", "", "생산연령(15~64세) 100명당 고령인구. SGIS 2024년"),
    ("평균나이", "실수", "세", "SGIS 2024년"),
    ("인구밀도", "실수", "명/㎢", "SGIS 2024년. SGIS 인구 기준이라 위 총인구와 분모가 다르다"),
    ("가구수", "정수", "가구", "SGIS 2024년 일반가구"),
    ("경로당수", "정수", "개소", "주소 기준 해당 읍면 소재 경로당 (2025-12-19 기준)"),
    ("노후경로당수", "정수", "개소", "건립연도 2000년 이전. 공식 기준이 아닌 프로젝트 임시 기준"),
    ("노후경로당비율", "실수", "%", "노후경로당수 / 경로당수 × 100"),
    ("고령인구1천명당_경로당수", "실수", "개소", "경로당수 / 고령인구 × 1,000"),
    ("중심위도", "실수", "도", "읍면 경계의 면적가중 중심점. EPSG:4326"),
    ("중심경도", "실수", "도", "읍면 경계의 면적가중 중심점. EPSG:4326"),
]


def load(name: str):
    return json.loads((DERIVED / name).read_text(encoding="utf-8"))


def _ring_centroid(ring: list) -> tuple[float, float, float]:
    """평면 근사 면적과 중심점. 읍면 크기에서는 경위도 평면 근사로 충분하다."""
    a = cx = cy = 0.0
    for (x0, y0), (x1, y1) in zip(ring, ring[1:]):
        cross = x0 * y1 - x1 * y0
        a += cross
        cx += (x0 + x1) * cross
        cy += (y0 + y1) * cross
    a /= 2
    return a, cx / (6 * a), cy / (6 * a)


def centroid(geom: dict) -> tuple[float, float]:
    """(위도, 경도). 구멍은 음의 면적으로 빼서 면적가중 평균한다."""
    polys = geom["coordinates"] if geom["type"] == "MultiPolygon" else [geom["coordinates"]]
    total = sx = sy = 0.0
    for poly in polys:
        for k, ring in enumerate(poly):
            a, x, y = _ring_centroid(ring)
            a = abs(a) if k == 0 else -abs(a)
            total += a
            sx += a * x
            sy += a * y
    return sy / total, sx / total


def build_rows() -> list[dict]:
    pop = [r for r in load("population.json")["records"]
           if r["period"] == ANALYSIS_PERIOD and r["level"] == "eupmyeon"]
    assert len(pop) == 8, f"{ANALYSIS_PERIOD} 읍면 인구가 8행이 아님: {len(pop)}"
    sgis = load("sgis_stats.json")["by_region"]
    seniors = load("senior_centers.json")["by_region"]
    shapes = {f["properties"]["region_code"]: f["geometry"]
              for f in load("eupmyeon.geojson")["features"]}

    rows = []
    for p in sorted(pop, key=lambda r: r["region_code"]):
        code = p["region_code"]
        lat, lon = centroid(shapes[code])
        rows.append({
            "종합점수": None,
            "행정구역코드": code,
            "시군구명": SIGUNGU,
            "읍면명": p["region_name"],
            "총인구": p["population"],
            "고령인구": p["elderly"],
            "고령화율": round(p["elderly"] / p["population"] * 100, 2),
            "노령화지수": sgis[code]["aging_index"],
            "노년부양비": sgis[code]["old_dependency"],
            "평균나이": sgis[code]["avg_age"],
            "인구밀도": sgis[code]["density_per_km2"],
            "가구수": sgis[code]["households"],
            "경로당수": seniors[code]["count"],
            "노후경로당수": seniors[code]["old_count"],
            "노후경로당비율": round(seniors[code]["old_count"] / seniors[code]["count"] * 100, 2),
            "고령인구1천명당_경로당수": round(seniors[code]["count"] / p["elderly"] * 1000, 2),
            "중심위도": round(lat, 6),
            "중심경도": round(lon, 6),
        })
    assert [c[0] for c in COLUMNS] == list(rows[0]), "COLUMNS 와 결과표 열이 어긋남"
    return rows


def method_rows() -> list[tuple[str, str]]:
    reg = {s["id"]: s for s in json.loads(REGISTRY.read_text(encoding="utf-8"))["sources"]}
    used = ["MOIS_AGE", "SGIS_ADM_STAT_2025", "DATAGOKR_GUNWI_SENIOR", "BOUNDARY_EUPMYEON"]
    excluded = [s for s in reg.values() if s["status"] == "분석제외"]

    rows = [
        ("■ 개요", ""),
        ("데이터셋", "대구광역시 군위군 읍면별 주거환경 취약지역 기초지표"),
        ("공간 단위", "읍·면 8곳 (군위읍, 소보·효령·부계·우보·의흥·산성·삼국유사면)"),
        ("기준 시점", "인구 2025년 12월 말 / SGIS 2024년 / 경로당 2025-12-19"),
        ("좌표계", "EPSG:4326 (WGS84)"),
        ("", ""),
        ("■ 산출식", ""),
        ("종합 점수", "계산하지 않는다. 지표 간 척도가 달라(건수·비율·지수) 표준화와 가중치 근거가 "
                     "없으므로 '종합점수' 열을 비워 두었다"),
        ("고령화율", "65세 이상 인구 ÷ 총인구 × 100"),
        ("고령인구1천명당_경로당수", "경로당수 ÷ 고령인구 × 1,000"),
        ("노후경로당수", "건립연도 < 2000. 공식 기준이 아니라 프로젝트 임시 기준"),
        ("", ""),
        ("■ 자료 출처", ""),
    ]
    for sid in used:
        s = reg[sid]
        rows.append((s["name"], f"{s['provider']} | 기준일 {s['reference_date']} | "
                                f"{s['url'] or 'URL 미확인'}"))
    rows += [
        ("", ""),
        ("■ 알아둘 점", ""),
        ("표본 수", "읍면이 8곳뿐이라 상관분석은 탐색적 수준이다. 유의성 해석에 주의"),
        ("두 가지 인구", "총인구·고령인구는 주민등록(2025-12), 노령화지수·인구밀도 등은 SGIS 등록센서스"
                      "(2024)다. 정의가 달라 값이 다르므로 한 지표 안에서 섞지 말 것"),
        ("경계 출처", "읍면 경계는 통계청 SGIS 행정구역 경계 2025년 2분기(기준일 2025-06-30). "
                    "2026-10-15 SGIS 원본과 8개 읍면 도형 일치 확인"),
        ("", ""),
        ("■ 분석에서 제외한 자료", ""),
    ]
    for s in excluded:
        reason = s["notes"].split("사유: ", 1)[-1].split(". 원본은", 1)[0]
        rows.append((s["name"], reason))
    rows += [
        ("", ""),
        ("■ 함께 제공하는 파일", ""),
        (f"{BASENAME}.csv", "결과표를 개방 포맷으로 (기계 판독용). 내용은 이 파일 1번 시트와 같다"),
        ("data/sources.json", "출처·기준일·확보 상태를 구조화한 출처 레지스트리"),
        ("tools/prepare.py · tools/export_excel.py", "원자료 → 정제 → 이 파일까지의 재현 코드"),
    ]
    return rows


HEADER_FONT = Font(bold=True, size=10, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
SECTION_FONT = Font(bold=True)
THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _sheet(ws, header: list[str], body: list[list], widths: list[float]) -> None:
    ws.append(header)
    for row in body:
        ws.append(row)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows():
        for cell in row:
            cell.border = BORDER
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"


def main() -> None:
    rows = build_rows()
    header = [c[0] for c in COLUMNS]
    OUT.mkdir(exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "기초지표"
    _sheet(ws, header, [[r[h] for h in header] for r in rows],
           [8, 13, 9, 11] + [13] * (len(header) - 4))
    for cell in ws["B"][1:]:
        cell.number_format = "@"          # 코드는 문자열로 둔다

    _sheet(wb.create_sheet("컬럼설명"), ["컬럼명", "자료형", "단위", "설명"],
           [list(c) for c in COLUMNS], [24, 9, 8, 84])

    ws3 = wb.create_sheet("산출방법")
    _sheet(ws3, ["구분", "내용"], [list(r) for r in method_rows()], [26, 112])
    for row in ws3.iter_rows(min_row=2):
        if str(row[0].value or "").startswith("■"):
            row[0].font = SECTION_FONT
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    xlsx = OUT / f"{BASENAME}.xlsx"
    wb.save(xlsx)
    with open(OUT / f"{BASENAME}.csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=header)
        w.writeheader()
        w.writerows(rows)

    print(f"  {xlsx.relative_to(ROOT)}  (읍면 {len(rows)}행 x {len(header)}열)")
    print(f"  {(OUT / f'{BASENAME}.csv').relative_to(ROOT)}")


if __name__ == "__main__":
    main()
